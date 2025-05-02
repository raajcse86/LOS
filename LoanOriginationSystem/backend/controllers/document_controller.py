import json
import datetime

from flask import jsonify, request, send_file, current_app
from services.document_service import DocumentService
from agents.compliance_validation import get_compliance_report
from werkzeug.utils import secure_filename
import os
import io
from bson.objectid import ObjectId
from bson.errors import InvalidId
import mimetypes
import traceback
from io import BytesIO
import pdfplumber
from io import BytesIO
from docx import Document

document_service = DocumentService()

def get_all_documents():
    """Get all documents."""
    documents = document_service.get_all_documents()
    return jsonify(documents)

def get_document(document_id):
    """Get a document by ID."""
    document = document_service.get_document_by_id(document_id)
    if not document:
        return jsonify({"error": "Document not found"}), 404
    return jsonify(document)

def update_document(document_id):
    """Update a document's status (approve/reject)."""
    if not request.json:
        return jsonify({"error": "No update data provided"}), 400
    
    # Get update data
    update_data = request.json
    
    # Validate status if provided
    status = update_data.get('status')
    if status and status not in ['Approved', 'Rejected', 'Pending Review']:
        return jsonify({"error": "Invalid status value"}), 400
    
    # Call service to update document
    success = document_service.update_document(document_id, update_data)
    if not success:
        return jsonify({"error": "Failed to update document"}), 500
    
    # Get updated document
    updated_document = document_service.get_document_by_id(document_id)
    return jsonify(updated_document)

def get_documents_by_loan(loan_id):
    """Get all documents for a specific loan."""
    documents = document_service.get_documents_by_loan_id(loan_id)
    return jsonify(documents)

def upload_document(loan_id):
    """Upload a document for a loan."""
    # Check if the request has the file part
    print(f"DEBUG - Upload document request for loan {loan_id}")
    print(f"DEBUG - Request method: {request.method}")
    print(f"DEBUG - Request content type: {request.content_type}")
    print(f"DEBUG - Request files: {request.files}")
    print(f"DEBUG - Request form: {request.form}")
    if 'file' not in request.files:
        print(f"DEBUG - No file part found in request.files: {list(request.files.keys())}")
        return jsonify({"error": "No file part"}), 400
    
    file = request.files['file']
    
    # Check if a file was selected
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    
    # Validate file type (allow .docx and .pdf)
    allowed_extensions = {'.docx', '.pdf'}
    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in allowed_extensions:
        return jsonify({"error": f"File type not allowed. Only {', '.join(allowed_extensions)} are supported"}), 400
    
    # Get document type from form data
    document_type = request.form.get('documentType')
    
    # Validate required fields
    if not document_type:
        return jsonify({"error": "documentType is required"}), 400
    
    # Upload document
    document = document_service.upload_document(loan_id, document_type, file)
    if not document:
        return jsonify({"error": "Failed to upload document"}), 500
    
    # Make sure document has string ID (already handled in service)
    # This is redundant but ensures we don't have ObjectId serialization issues
    if document and '_id' in document and not isinstance(document['_id'], str):
        document['_id'] = str(document['_id'])
    
    return jsonify(document), 201

def download_document(document_id):
    """Download a document."""
    current_app.logger.info(f"Download request received for document_id: {document_id}")
    
    try:
        result = document_service.download_document(document_id)
        if not result:
            current_app.logger.warning(f"Document not found: {document_id}")
            return jsonify({"error": "Document not found"}), 404
        
        current_app.logger.info(f"Sending file: {result['filename']}")
        return send_file(
            io.BytesIO(result['file_data']),
            download_name=result['filename'],
            mimetype=result['mime_type']
        )
    except Exception as e:
        current_app.logger.error(f"Error downloading document: {str(e)}", exc_info=True)
        return jsonify({"error": "Failed to download document"}), 500

def extract_document_content(document_id):
    """Extract content from a document."""
    result = document_service.extract_document_content(document_id)
    if not result:
        return jsonify({"error": "Failed to extract document content"}), 500
    
    return jsonify(result)

def validate_document(document_id):
    """Validate compliance for a specific document."""
    try:
        # Get the document binary data
        result = document_service.download_document(document_id)
        if not result:
            return jsonify({"error": "Document not found"}), 404

        file_data = result['file_data']
        mime_type = result['mime_type']
        filename = result.get('filename', '')
        
        current_app.logger.info(f"Processing document: {filename} with mime type: {mime_type}")
        
        # Extract text content based on file type
        document_content = ""
        
        try:
            if mime_type == 'application/pdf' or filename.lower().endswith('.pdf'):
                # Extract text from PDF using pdfplumber
                
                
                current_app.logger.info("Processing PDF document")
                with BytesIO(file_data) as pdf_buffer:
                    with pdfplumber.open(pdf_buffer) as pdf:
                        for page in pdf.pages:
                            extracted_text = page.extract_text()
                            if extracted_text:
                                document_content += extracted_text + "\n"
            
            elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' or filename.lower().endswith('.docx'):
                # Extract text from DOCX using python-docx
                
                
                current_app.logger.info("Processing DOCX document")
                doc = Document(BytesIO(file_data))
                
                # Extract text from paragraphs
                paragraphs_text = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
                
                # Extract text from tables
                tables_text = []
                for table in doc.tables:
                    for row in table.rows:
                        row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            tables_text.append(row_text)
                
                # Combine all text with proper formatting
                document_content = "\n".join(paragraphs_text)
                if tables_text:
                    document_content += "\n\nTables Content:\n" + "\n".join(tables_text)
            
            else:
                current_app.logger.error(f"Unsupported file type: {mime_type} for file {filename}")
                return jsonify({
                    "error": "Unsupported file type",
                    "message": "Only PDF and DOCX files are supported"
                }), 400

        except Exception as e:
            current_app.logger.error(f"Error extracting text from document: {str(e)}", exc_info=True)
            return jsonify({
                "error": "Text extraction failed",
                "message": f"Failed to extract text from the document: {str(e)}"
            }), 500

        # Validate extracted content
        if not document_content.strip():
            current_app.logger.warning(f"No text content extracted from document: {filename}")
            return jsonify({
                "error": "Empty content",
                "message": "No text content could be extracted from the document"
            }), 400

        current_app.logger.info(f"Successfully extracted {len(document_content)} characters from document")

        # Define document rules
        document_rules = """Purpose: Validate current employment status and employer relationship.
Checklist:
•	Must be on company letterhead with logo and official contact details.
•	Must include:
o	Date of issue (within the past 30 days).
o	Full name of employee.
o	Employment status (Full-time/Part-time/Contract).
o	Date of joining.
o	Job title.
o	Salary or hourly wage and pay frequency.
•	Must be signed by authorized HR personnel or supervisor.
•	Should include a valid contact for verification.
•	Should not contain discrepancies from submitted payslips or application.
"""

        # Get compliance report
        try:
            current_app.logger.info("Generating compliance report")
            validation_response = get_compliance_report(document_content, document_rules)
            
            result = json.loads(validation_response)
            
            # Add metadata to response
            result['metadata'] = {
                'filename': filename,
                'document_type': 'PDF' if mime_type == 'application/pdf' or filename.lower().endswith('.pdf') else 'DOCX',
                'content_length': len(document_content),
                'processed_at': datetime.datetime.utcnow().isoformat()
            }
            
            return jsonify(result)
            
        except json.JSONDecodeError as e:
            current_app.logger.error(f"Invalid compliance report format: {str(e)}")
            return jsonify({
                "error": "Invalid compliance report format",
                "message": "The compliance report could not be processed"
            }), 500
        except Exception as e:
            current_app.logger.error(f"Error generating compliance report: {str(e)}", exc_info=True)
            return jsonify({
                "error": "Compliance validation failed",
                "message": f"Failed to generate compliance report: {str(e)}"
            }), 500

    except Exception as e:
        current_app.logger.error(f"Error in validate_document: {str(e)}", exc_info=True)
        return jsonify({
            "error": "Processing failed",
            "message": f"Failed to process document: {str(e)}"
        }), 500

def check_compliance_score(document_id):
    """Check compliance score for a document."""
    result = document_service.check_compliance_score(document_id)
    if not result:
        return jsonify({"error": "Failed to check compliance score"}), 500
    
    return jsonify(result)

def delete_document(document_id):
    """Delete a document."""
    success = document_service.delete_document(document_id)
    if not success:
        return jsonify({"error": "Failed to delete document"}), 500
    
    return jsonify({"message": "Document deleted successfully"})
